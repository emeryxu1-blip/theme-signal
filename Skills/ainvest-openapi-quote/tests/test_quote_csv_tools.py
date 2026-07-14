import csv
import json
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
import sys


SKILL_ROOT = Path(__file__).resolve().parent.parent
if str(SKILL_ROOT) not in sys.path:
    sys.path.insert(0, str(SKILL_ROOT))

from scripts import build_quote_csvs
from scripts import export_indicators
from scripts import find_quote_params

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None


TMP_ROOT = Path(tempfile.gettempdir()) / "ainvest-openapi-quote-test-workspaces"


def workspace_case_dir(name):
    TMP_ROOT.mkdir(exist_ok=True)
    path = TMP_ROOT / name
    path.mkdir(exist_ok=True)
    return path


class FakeExportResponse:
    def __init__(self, status_code=200, content=b"workbook-bytes", text=""):
        self.status_code = status_code
        self.content = content
        self.text = text


class ExportIndicatorsTest(unittest.TestCase):
    def test_default_export_and_build_paths_are_inside_skill(self):
        expected = SKILL_ROOT / "references" / "generated" / "export_metric_meta_new.xlsx"

        self.assertEqual(export_indicators.DEFAULT_OUTPUT, expected)
        self.assertEqual(build_quote_csvs.DEFAULT_INPUT, expected)

    def test_default_http_post_uses_stdlib_compatible_response_shape(self):
        captured = {}

        class FakeHandle:
            status = 200

            def read(self):
                return b"xlsx"

        def fake_urlopen(request, timeout):
            captured["url"] = request.full_url
            captured["data"] = request.data
            captured["content_type"] = request.get_header("Content-type")
            captured["timeout"] = timeout
            return FakeHandle()

        response = export_indicators.default_http_post(
            export_indicators.URL,
            headers={"Content-Type": "application/json"},
            json={"tenant_id": "ainvest"},
            timeout=9,
            urlopen=fake_urlopen,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, b"xlsx")
        self.assertEqual(response.text, "")
        self.assertEqual(captured["url"], export_indicators.URL)
        self.assertEqual(captured["data"], b'{"tenant_id": "ainvest"}')
        self.assertEqual(captured["content_type"], "application/json")
        self.assertEqual(captured["timeout"], 9)

    def test_local_data_status_reports_missing_workbook(self):
        tmp_path = workspace_case_dir("export_indicators_missing_status")
        workbook_path = tmp_path / "missing.xlsx"

        status = export_indicators.local_data_status(
            workbook_path,
            now=datetime(2026, 5, 26, 12, 0, tzinfo=timezone.utc),
        )

        self.assertEqual(status["path"], str(workbook_path))
        self.assertEqual(status["exists"], False)
        self.assertEqual(status["age_seconds"], None)
        self.assertEqual(status["age_text"], "not available")

    def test_local_data_status_reports_workbook_age(self):
        tmp_path = workspace_case_dir("export_indicators_existing_status")
        workbook_path = tmp_path / "export_metric_meta_new.xlsx"
        workbook_path.write_bytes(b"xlsx")
        mtime = datetime(2026, 5, 24, 9, 30, tzinfo=timezone.utc)
        timestamp = mtime.timestamp()
        import os

        os.utime(workbook_path, (timestamp, timestamp))

        status = export_indicators.local_data_status(
            workbook_path,
            now=datetime(2026, 5, 26, 12, 0, tzinfo=timezone.utc),
        )

        self.assertEqual(status["exists"], True)
        self.assertEqual(status["size_bytes"], 4)
        self.assertEqual(status["updated_at"], "2026-05-24T09:30:00+00:00")
        self.assertEqual(status["age_seconds"], int(timedelta(days=2, hours=2, minutes=30).total_seconds()))
        self.assertEqual(status["age_text"], "2 days 2 hours")

    def test_update_indicators_can_rebuild_csvs_after_export(self):
        tmp_path = workspace_case_dir("export_indicators_rebuild")
        workbook_path = tmp_path / "export_metric_meta_new.xlsx"
        calls = []

        def fake_post(url, headers, json, timeout):
            return FakeExportResponse(content=b"xlsx")

        def fake_run(command, cwd, check, text, capture_output):
            calls.append(
                {
                    "command": command,
                    "cwd": cwd,
                    "check": check,
                    "text": text,
                    "capture_output": capture_output,
                }
            )
            return export_indicators.RebuildResult(returncode=0, stdout="built", stderr="")

        path, byte_count, rebuild_result = export_indicators.update_indicators(
            output_path=workbook_path,
            http_post=fake_post,
            rebuild_csvs=True,
            command_runner=fake_run,
        )

        self.assertEqual(path, workbook_path)
        self.assertEqual(byte_count, 4)
        self.assertEqual(rebuild_result.stdout, "built")
        self.assertEqual(calls[0]["cwd"], export_indicators.SKILL_ROOT)
        self.assertEqual(calls[0]["check"], False)
        self.assertIn(str(workbook_path), calls[0]["command"])
        self.assertIn("--input", calls[0]["command"])

    def test_export_indicators_writes_workbook_and_uses_tangram_params(self):
        tmp_path = workspace_case_dir("export_indicators_success")
        output_path = tmp_path / "export_metric_meta_new.xlsx"
        calls = []

        def fake_post(url, headers, json, timeout):
            calls.append({"url": url, "headers": headers, "json": json, "timeout": timeout})
            return FakeExportResponse(content=b"xlsx")

        path, byte_count, params = export_indicators.export_indicators(
            output_path=output_path,
            keyword="price",
            page_size=50,
            timeout=12,
            http_post=fake_post,
        )

        self.assertEqual(path, output_path)
        self.assertEqual(byte_count, 4)
        self.assertEqual(output_path.read_bytes(), b"xlsx")
        self.assertEqual(calls[0]["url"], export_indicators.URL)
        self.assertEqual(calls[0]["headers"]["Content-Type"], "application/json")
        self.assertEqual(calls[0]["json"]["tenant_id"], "ainvest")
        self.assertEqual(calls[0]["json"]["keyword"], "price")
        self.assertEqual(calls[0]["json"]["page_size"], 50)
        self.assertEqual(calls[0]["timeout"], 12)
        self.assertEqual(params["support_api"], True)

    def test_export_indicators_raises_on_tangram_error(self):
        tmp_path = workspace_case_dir("export_indicators_failure")
        output_path = tmp_path / "export_metric_meta_new.xlsx"

        def fake_post(url, headers, json, timeout):
            return FakeExportResponse(status_code=500, text="boom")

        with self.assertRaisesRegex(RuntimeError, "Tangram export failed"):
            export_indicators.export_indicators(output_path=output_path, http_post=fake_post)

        self.assertFalse(output_path.exists())


@unittest.skipUnless(Workbook, "openpyxl is required for Excel builder tests")
class QuoteCsvToolsTest(unittest.TestCase):
    def write_workbook(self, path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet.merge_cells("A1:X1")
        sheet["A1"] = "Tangram template"
        sheet.append([])
        sheet.append([])
        sheet.append(
            [
                "*指标编码",
                "*来源代码",
                "*标准编码",
                "*指标名称",
                "英文名称",
                "*指标描述",
                "*指标来源",
                "*数据建模团队",
                "*数据类型",
                "*支持周期",
                "*默认周期",
                "*业务领域UIDs",
                "所属分类UIDs",
                "*指标状态",
                "计算公式",
                "指标业务背景",
                "*指标负责人邮箱",
                "来源指标代码别名",
                "*数据更新时间",
                "显示值类型",
                "指标单位",
                "自定义取数参数",
                "扩展属性",
                "支持API",
            ]
        )
        sheet.append(
            [
                "F10-volume",
                "13",
                "volume_standard",
                "成交量",
                "Volume",
                "成交量说明",
                "F10",
                "F10_TEAM",
                "BIG_DECIMAL",
                "SNAPSHOT",
                "SNAPSHOT",
                "",
                "",
                "40",
                "",
                "",
                "",
                "",
                "",
                "normal",
                "股",
                "[]",
                (
                    '[{"required":"1","name":"time_period","example":"DAY_1",'
                    '"query_type":"enum","default_value":"DAY_1",'
                    '"enum_options":[{"value":"DAY_1","label":"1日"}],'
                    '"children":null}]'
                    '[AUTO_ATTRS_FROM_ID_DICT]'
                    '[{"name":"time_period","default":"intraday"},'
                    '{"name":"trade_class","default":"intraday"}]'
                ),
                "1",
            ]
        )
        sheet.append(
            [
                "F10-stoch",
                "",
                "stochastic_rsi",
                "随机RSI",
                "Stochastic RSI",
                "技术指标说明",
                "F10",
                "F10_TEAM",
                "BIG_DECIMAL",
                "DAY_1",
                "DAY_1",
                "",
                "",
                "40",
                "",
                "",
                "",
                "",
                "",
                "ratio",
                "%",
                "",
                (
                    '[{"required":"1","name":"tech_param","query_type":"array",'
                    '"default_value":null,"enum_options":null,'
                    '"children":[{"required":"1","name":"indicator_component",'
                    '"default_value":"K","enum_options":[{"value":"K","label":"K线"}],'
                    '"children":null}]}]'
                ),
                "1",
            ]
        )
        workbook.save(path)

    def write_new_tangram_workbook(self, path, attrs_text, access_guide, entity_type, exchange, sortable="1"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Tangram 导入模版20260107-v3.0(*必填)"] * 30)
        sheet.append(["导入模版请勿调整-导入时 请注意导入的环境"] * 30)
        sheet.append([])
        sheet.append(
            [
                "*指标代码(废弃)",
                "*指标代码",
                "*IndexAPI代码",
                "*来源代码",
                "*指标名称",
                "英文名称",
                "*指标描述",
                "*指标来源",
                "*数据建模团队",
                "*数据类型",
                "所属分类UIDs",
                "所属分类(仅展示)",
                "指标公式",
                "业务背景",
                "*负责人",
                "显示值类型",
                "指标单位",
                "自定义取数参数",
                "扩展属性",
                "支持过滤",
                "支持排序",
                "支持搜索",
                "缓存策略",
                "指标标签",
                "默认代码选择器",
                "接入指南标签",
                "证券实体类型",
                "交易所",
                "血缘",
                "返回值多语言",
            ]
        )
        sheet.append(
            [
                "old_code",
                "display_code",
                "idx_price_change",
                "source_should_not_be_used",
                "涨幅",
                "Price Change",
                "涨幅说明",
                "AINVEST",
                "AINVEST_TEAM",
                "BIG_DECIMAL",
                "",
                "",
                "",
                "",
                "owner@example.com",
                "ratio",
                "%",
                "",
                attrs_text,
                "0",
                sortable,
                "1",
                "0",
                "0",
                "",
                access_guide,
                entity_type,
                exchange,
                "[]",
                "",
            ]
        )
        workbook.save(path)

    def test_build_outputs_id_dict_and_request_lookup_from_excel(self):
        tmp_path = workspace_case_dir("build_case")
        workbook_path = tmp_path / "source.xlsx"
        out_dir = tmp_path / "generated"
        self.write_workbook(workbook_path)

        result = build_quote_csvs.build_csvs(workbook_path, out_dir, legacy_id_dict_path="")

        self.assertEqual(result.rows_written, 2)
        with (out_dir / "id_dict_quote.csv").open(encoding="utf-8-sig") as handle:
            id_rows = list(csv.DictReader(handle))
        with (out_dir / "quote_request_lookup.csv").open(encoding="utf-8-sig") as handle:
            lookup_rows = list(csv.DictReader(handle))
        self.assertEqual(id_rows[0]["id"], "13")
        self.assertEqual(id_rows[0]["periods"], "day_1")
        self.assertEqual(id_rows[0]["attrs"], "time_period;trade_class:intraday")
        self.assertEqual(lookup_rows[0]["indicator_id"], "13")
        self.assertEqual(lookup_rows[0]["endpoint"], "series")
        self.assertEqual(lookup_rows[0]["category"], "security")

        attrs = json.loads(lookup_rows[0]["attrs_json"])
        self.assertEqual(attrs["time_period"]["default"], "day_1")
        self.assertEqual(attrs["time_period"]["enum_options"][0]["value"], "day_1")
        self.assertEqual(attrs["trade_class"]["default"], "intraday")

        tech_attrs = json.loads(lookup_rows[1]["attrs_json"])
        self.assertEqual(lookup_rows[1]["indicator_id"], "stochastic_rsi")
        self.assertEqual(lookup_rows[1]["endpoint"], "series")
        self.assertEqual(tech_attrs["tech_param"]["children"][0]["name"], "indicator_component")

    def test_build_outputs_uses_legacy_attrs_as_first_priority(self):
        tmp_path = workspace_case_dir("legacy_attr_case")
        workbook_path = tmp_path / "source.xlsx"
        out_dir = tmp_path / "generated"
        legacy_path = tmp_path / "id_dict.md"
        self.write_workbook(workbook_path)
        legacy_path.write_text(
            "\n".join(
                [
                    "| id | name | desc | attrs |",
                    "| --- | --- | --- | --- |",
                    "| 13 | volume | volume | event_id 必选time_period 必选 day_1 |",
                ]
            ),
            encoding="utf-8",
        )

        build_quote_csvs.build_csvs(workbook_path, out_dir, legacy_path)

        with (out_dir / "id_dict_quote.csv").open(encoding="utf-8-sig") as handle:
            id_rows = list(csv.DictReader(handle))
        with (out_dir / "quote_request_lookup.csv").open(encoding="utf-8-sig") as handle:
            lookup_rows = list(csv.DictReader(handle))
        attrs_json = json.loads(lookup_rows[0]["attrs_json"])

        self.assertIn("event_id", id_rows[0]["attrs"].split(";"))
        self.assertIn("time_period", id_rows[0]["attrs"].split(";"))
        self.assertEqual(attrs_json["event_id"]["source"], "legacy")
        self.assertEqual(attrs_json["time_period"]["source"], "legacy")
        self.assertEqual(attrs_json["time_period"]["default"], "day_1")

    def test_build_outputs_uses_legacy_time_range_as_first_priority(self):
        tmp_path = workspace_case_dir("legacy_time_range_case")
        workbook_path = tmp_path / "source.xlsx"
        out_dir = tmp_path / "generated"
        legacy_path = tmp_path / "id_dict.md"
        self.write_workbook(workbook_path)
        legacy_path.write_text(
            "\n".join(
                [
                    "| id | 名称 | time_period | time_range支持的方式 |",
                    "| --- | --- | --- | --- |",
                    "| stochastic_rsi | 随机RSI | day_1 | begin_end |",
                    "| volume_call | 期权统计--看涨成交量 | min_1 | end_count {end_time:0,count:1、2}  trade_date = 0 |",
                ]
            ),
            encoding="utf-8",
        )

        build_quote_csvs.build_csvs(workbook_path, out_dir, legacy_path)

        with (out_dir / "quote_request_lookup.csv").open(encoding="utf-8-sig") as handle:
            lookup_rows = {row["indicator_id"]: row for row in csv.DictReader(handle)}
        time_range = json.loads(lookup_rows["stochastic_rsi"]["time_range_json"])
        special_rule = build_quote_csvs.load_legacy_time_ranges(legacy_path)["volume_call"]
        special_time_range = json.loads(build_quote_csvs.build_time_range_json("series", "min_1", special_rule))
        fixed_begin_time = 1777939200000
        begin_end_time_range = json.loads(
            build_quote_csvs.build_time_range_json(
                "series",
                "day_1",
                {"time_period": "day_1", "time_range": "begin_end"},
                fixed_begin_time,
            )
        )

        self.assertEqual(time_range["type"], "begin_end")
        self.assertGreater(time_range["begin_time"], 0)
        self.assertEqual(time_range["end_time"], 0)
        self.assertEqual(special_time_range, {"type": "trade_date", "trade_date": 0, "time_period": "min_1"})
        self.assertEqual(
            begin_end_time_range,
            {"type": "begin_end", "begin_time": fixed_begin_time, "end_time": 0, "time_period": "day_1"},
        )

    def test_build_outputs_use_new_tangram_columns(self):
        tmp_path = workspace_case_dir("new_tangram_columns")
        workbook_path = tmp_path / "source.xlsx"
        out_dir = tmp_path / "generated"
        attrs = json.dumps(
            [
                {
                    "required": "1",
                    "name": "time_period",
                    "example": "day_1",
                    "query_type": "enum",
                    "default_value": "day_1",
                    "enum_options": [
                        {"value": "day_1", "label": "日"},
                        {"value": "min_5", "label": "5分钟"},
                    ],
                    "children": None,
                }
            ],
            ensure_ascii=False,
        )
        self.write_new_tangram_workbook(
            workbook_path,
            attrs,
            "STANDARD_SERIES_TIME_RANGE_BEGIN_END",
            "security",
            "UUS",
        )

        result = build_quote_csvs.build_csvs(workbook_path, out_dir, legacy_id_dict_path="", category_map_path="")

        self.assertEqual(result.rows_written, 1)
        with (out_dir / "id_dict_quote.csv").open(encoding="utf-8-sig") as handle:
            id_row = next(csv.DictReader(handle))
        with (out_dir / "quote_request_lookup.csv").open(encoding="utf-8-sig") as handle:
            lookup_row = next(csv.DictReader(handle))
        attrs_json = json.loads(lookup_row["attrs_json"])
        time_range = json.loads(lookup_row["time_range_json"])

        self.assertEqual(id_row["id"], "idx_price_change")
        self.assertEqual(id_row["periods"], "day_1;min_5")
        self.assertEqual(id_row["sortable"], "TRUE")
        self.assertEqual(lookup_row["indicator_id"], "idx_price_change")
        self.assertNotIn("source_should_not_be_used", lookup_row["query_key"])
        self.assertEqual(lookup_row["endpoint"], "series")
        self.assertEqual(lookup_row["category"], "security")
        self.assertEqual(lookup_row["symbol_type"], "market_code")
        self.assertEqual(lookup_row["exchange"], "UUS")
        self.assertEqual(lookup_row["access_guide"], "STANDARD_SERIES_TIME_RANGE_BEGIN_END")
        self.assertEqual(attrs_json["time_period"]["default"], "day_1")
        self.assertEqual(lookup_row["period_values_json"], '["day_1","min_5"]')
        self.assertEqual(time_range["type"], "begin_end")
        self.assertEqual(time_range["time_period"], "day_1")

    def test_build_outputs_fall_back_to_router_periods_when_attrs_have_no_time_period(self):
        tmp_path = workspace_case_dir("new_tangram_router_periods")
        workbook_path = tmp_path / "source.xlsx"
        out_dir = tmp_path / "generated"
        router_path = tmp_path / "id_router.yaml"
        self.write_new_tangram_workbook(
            workbook_path,
            "[]",
            "STANDARD_SERIES_TIME_RANGE_END_COUNT",
            "market_env",
            "UBAX",
            sortable="0",
        )
        router_path.write_text(
            "\n".join(
                [
                    "router:",
                    "  - service: test",
                    "    markets: [UBAX]",
                    "    routes:",
                    "      - api: /history",
                    "        ids: [idx_price_change]",
                    "        periods: [day_1, week_1]",
                    "        sortable: false",
                ]
            ),
            encoding="utf-8",
        )

        build_quote_csvs.build_csvs(
            workbook_path,
            out_dir,
            legacy_id_dict_path="",
            category_map_path="",
            router_path=router_path,
        )

        with (out_dir / "quote_request_lookup.csv").open(encoding="utf-8-sig") as handle:
            lookup_row = next(csv.DictReader(handle))
        time_range = json.loads(lookup_row["time_range_json"])

        self.assertEqual(lookup_row["periods"], "day_1;week_1")
        self.assertEqual(lookup_row["category"], "market_env")
        self.assertEqual(lookup_row["symbol_type"], "")
        self.assertEqual(lookup_row["exchange"], "UBAX")
        self.assertEqual(time_range["type"], "end_count")
        self.assertEqual(time_range["time_period"], "day_1")

class QuoteLookupTest(unittest.TestCase):
    def test_find_quote_params_searches_generated_lookup(self):
        tmp_path = workspace_case_dir("find_case")
        lookup_path = tmp_path / "quote_request_lookup.csv"
        with lookup_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=find_quote_params.LOOKUP_FIELDS)
            writer.writeheader()
            writer.writerow(
                {
                    "query_key": "13|volume_standard|成交量|Volume",
                    "category": "security",
                    "scenario": "metric_lookup",
                    "endpoint": "snapshot",
                    "symbol_type": "market_code",
                    "indicator_id": "13",
                    "req_unique_id": "13",
                    "attrs_json": "{}",
                    "time_range_json": "{}",
                    "template": "stock-detail.json",
                    "periods": "snapshot",
                    "metric_name": "成交量",
                    "english_name": "Volume",
                    "description": "成交量说明",
                    "aliases": "",
                    "is_interval_metric": "false",
                    "period_values_json": "[]",
                }
            )

        matches = find_quote_params.find_matches(lookup_path, "volume")

        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0]["indicator_id"], "13")

    def test_interval_aliases_do_not_include_time_words(self):
        aliases = build_quote_csvs.supplemental_aliases(
            "inr-price_change_ratio_pct-sum",
            "区间涨幅",
        )

        self.assertIn("区间收盘价涨幅", aliases)
        self.assertIn("区间涨幅", aliases)
        self.assertFalse(any("分钟" in alias or "MIN_" in alias.upper() for alias in aliases))

    def test_interval_query_prefers_interval_metric_but_plain_5m_prefers_fixed_metric(self):
        tmp_path = workspace_case_dir("rank_case")
        lookup_path = tmp_path / "quote_request_lookup.csv"
        fields = find_quote_params.LOOKUP_FIELDS
        with lookup_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerow(
                {
                    "query_key": "3934664|price_change_ratio_pct_5m|5分钟涨幅|5 Minute Price Change %",
                    "category": "security",
                    "scenario": "metric_lookup",
                    "endpoint": "snapshot",
                    "symbol_type": "market_code",
                    "indicator_id": "3934664",
                    "req_unique_id": "3934664",
                    "attrs_json": "{}",
                    "time_range_json": "{}",
                    "template": "stock-detail.json",
                    "periods": "snapshot",
                    "metric_name": "5分钟涨幅",
                    "english_name": "5 Minute Price Change %",
                    "description": "5分钟涨幅",
                    "aliases": "5分钟涨速|5分钟涨幅",
                    "is_interval_metric": "false",
                    "period_values_json": "[]",
                }
            )
            writer.writerow(
                {
                    "query_key": "inr-open_price_change_ratio_pct_diff-sum|区间涨幅差",
                    "category": "security",
                    "scenario": "metric_lookup",
                    "endpoint": "snapshot",
                    "symbol_type": "market_code",
                    "indicator_id": "inr-open_price_change_ratio_pct_diff-sum",
                    "req_unique_id": "inr-open_price_change_ratio_pct_diff-sum",
                    "attrs_json": "{}",
                    "time_range_json": "{}",
                    "template": "stock-detail.json",
                    "periods": "snapshot",
                    "metric_name": "区间涨幅差",
                    "english_name": "Difference of Price Change Ratio",
                    "description": "区间涨幅差",
                    "aliases": "区间涨幅差",
                    "is_interval_metric": "true",
                    "period_values_json": '["MIN_1","MIN_5"]',
                }
            )
            writer.writerow(
                {
                    "query_key": "inr-price_change_ratio_pct-sum|区间涨幅|Price Change % (Period)",
                    "category": "security",
                    "scenario": "metric_lookup",
                    "endpoint": "snapshot",
                    "symbol_type": "market_code",
                    "indicator_id": "inr-price_change_ratio_pct-sum",
                    "req_unique_id": "inr-price_change_ratio_pct-sum",
                    "attrs_json": json.dumps(
                        {
                            "time_period": {
                                "enum_options": [
                                    {"value": "MIN_1", "label": "1分钟"},
                                    {"value": "MIN_5", "label": "5分钟"},
                                ]
                            },
                            "period_type": {},
                        },
                        ensure_ascii=False,
                    ),
                    "time_range_json": "{}",
                    "template": "stock-detail.json",
                    "periods": "snapshot",
                    "metric_name": "区间涨幅",
                    "english_name": "Price Change % (Period)",
                    "description": "区间涨幅",
                    "aliases": "区间收盘价涨幅|区间涨幅",
                    "is_interval_metric": "true",
                    "period_values_json": '["MIN_1","MIN_5"]',
                }
            )

        interval_matches = find_quote_params.find_matches(lookup_path, "5分钟区间涨幅")
        plain_matches = find_quote_params.find_matches(lookup_path, "5分钟涨幅")

        self.assertEqual(interval_matches[0]["indicator_id"], "inr-price_change_ratio_pct-sum")
        self.assertEqual(plain_matches[0]["indicator_id"], "3934664")

    def test_market_env_rows_do_not_require_market_code_and_do_not_mix_with_security(self):
        tmp_path = workspace_case_dir("market_env_case")
        lookup_path = tmp_path / "quote_request_lookup.csv"
        with lookup_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=find_quote_params.LOOKUP_FIELDS)
            writer.writeheader()
            writer.writerow(
                {
                    "query_key": "ext_metric_altcoin_season_index|山寨币指数",
                    "category": "market_env",
                    "scenario": "metric_lookup",
                    "endpoint": "snapshot",
                    "symbol_type": "",
                    "indicator_id": "ext_metric_altcoin_season_index",
                    "req_unique_id": "ext_metric_altcoin_season_index",
                    "attrs_json": "{}",
                    "time_range_json": "{}",
                    "template": "market-env-snapshot.json",
                    "periods": "snapshot",
                    "metric_name": "山寨币指数",
                    "english_name": "",
                    "description": "山寨币指数",
                    "aliases": "",
                    "is_interval_metric": "false",
                    "period_values_json": "[]",
                }
            )
            writer.writerow(
                {
                    "query_key": "ext_metric_security_class|证券类型",
                    "category": "security",
                    "scenario": "metric_lookup",
                    "endpoint": "snapshot",
                    "symbol_type": "market_code",
                    "indicator_id": "ext_metric_security_class",
                    "req_unique_id": "ext_metric_security_class",
                    "attrs_json": "{}",
                    "time_range_json": "{}",
                    "template": "stock-detail.json",
                    "periods": "snapshot",
                    "metric_name": "证券类型",
                    "english_name": "",
                    "description": "证券类型",
                    "aliases": "",
                    "is_interval_metric": "false",
                    "period_values_json": "[]",
                }
            )

        matches = find_quote_params.find_matches(lookup_path, "山寨币指数")

        self.assertEqual(matches[0]["indicator_id"], "ext_metric_altcoin_season_index")
        self.assertEqual(matches[0]["category"], "market_env")
        self.assertEqual(matches[0]["symbol_type"], "")
        self.assertEqual(matches[0]["template"], "market-env-snapshot.json")
        self.assertTrue(all(row["category"] == "market_env" for row in matches))


if __name__ == "__main__":
    unittest.main()
